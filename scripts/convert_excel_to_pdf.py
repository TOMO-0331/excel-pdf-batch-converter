from __future__ import annotations

import argparse
import json
import logging
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, range_boundaries


LOGGER = logging.getLogger("excel-pdf-converter")


@dataclass(frozen=True)
class DeleteRule:
    sheet: str
    range_text: str
    mode: str


ROW_RANGE_RE = re.compile(r"^(?P<start>[1-9][0-9]*):(?P<end>[1-9][0-9]*)$")
COL_RANGE_RE = re.compile(r"^(?P<start>[A-Z]{1,3}):(?P<end>[A-Z]{1,3})$")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Edit Excel files and convert them to PDF with LibreOffice.")
    parser.add_argument("--job-dir", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--zip-path", required=True, type=Path)
    return parser.parse_args()


def read_config(job_dir: Path) -> dict:
    config_path = job_dir / "job-config.json"
    if not config_path.exists():
        raise FileNotFoundError(f"Job config not found: {config_path}")
    with config_path.open("r", encoding="utf-8") as file:
        return json.load(file)


def build_rules(raw_rules: list[dict]) -> list[DeleteRule]:
    rules: list[DeleteRule] = []
    for raw in raw_rules:
        sheet = str(raw.get("sheet", "")).strip()
        range_text = str(raw.get("range", "")).strip().upper()
        mode = str(raw.get("mode", "")).strip()
        if not sheet or not range_text or mode not in {"clear_values", "delete_rows", "delete_columns"}:
            raise ValueError(f"Invalid rule: {raw}")
        rules.append(DeleteRule(sheet=sheet, range_text=range_text, mode=mode))
    return rules


def apply_clear_values(sheet, range_text: str) -> None:
    min_col, min_row, max_col, max_row = range_boundaries(range_text)
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.value = None


def apply_delete_rows(sheet, range_text: str) -> None:
    row_match = ROW_RANGE_RE.match(range_text)
    if row_match:
        start = int(row_match.group("start"))
        end = int(row_match.group("end"))
    else:
        _, start, _, end = range_boundaries(range_text)
    if end < start:
        raise ValueError(f"Invalid row range: {range_text}")
    sheet.delete_rows(start, end - start + 1)


def apply_delete_columns(sheet, range_text: str) -> None:
    col_match = COL_RANGE_RE.match(range_text)
    if col_match:
        start = column_index_from_string(col_match.group("start"))
        end = column_index_from_string(col_match.group("end"))
    else:
        start, _, end, _ = range_boundaries(range_text)
    if end < start:
        raise ValueError(f"Invalid column range: {range_text}")
    sheet.delete_cols(start, end - start + 1)


def apply_rules(workbook_path: Path, output_path: Path, rules: list[DeleteRule]) -> None:
    workbook = load_workbook(workbook_path)
    for rule in rules:
        if rule.sheet not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {rule.sheet}")
        worksheet = workbook[rule.sheet]
        LOGGER.info("Applying %s to %s!%s", rule.mode, rule.sheet, rule.range_text)
        if rule.mode == "clear_values":
            apply_clear_values(worksheet, rule.range_text)
        elif rule.mode == "delete_rows":
            apply_delete_rows(worksheet, rule.range_text)
        elif rule.mode == "delete_columns":
            apply_delete_columns(worksheet, rule.range_text)
        else:
            raise ValueError(f"Unsupported mode: {rule.mode}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def convert_to_pdf(excel_path: Path, pdf_dir: Path) -> Path:
    pdf_dir.mkdir(parents=True, exist_ok=True)
    command = [
        "soffice",
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        "--convert-to",
        "pdf",
        "--outdir",
        str(pdf_dir),
        str(excel_path),
    ]
    LOGGER.info("Running LibreOffice for %s", excel_path.name)
    result = subprocess.run(command, check=False, capture_output=True, text=True)
    if result.stdout:
        LOGGER.info(result.stdout.strip())
    if result.stderr:
        LOGGER.warning(result.stderr.strip())
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice failed for {excel_path.name}: {result.stderr or result.stdout}")

    pdf_path = pdf_dir / f"{excel_path.stem}.pdf"
    if not pdf_path.exists():
        candidates = sorted(pdf_dir.glob("*.pdf"))
        if len(candidates) == 1:
            return candidates[0]
        raise FileNotFoundError(f"PDF was not created for {excel_path.name}")
    return pdf_path


def create_zip(source_dir: Path, zip_path: Path) -> None:
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(source_dir.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(source_dir))


def process_job(job_dir: Path, output_dir: Path, zip_path: Path) -> int:
    config = read_config(job_dir)
    input_dir = job_dir / "input"
    edited_dir = output_dir / "edited"
    pdf_dir = output_dir / "pdf"
    logs: list[dict] = []
    failures = 0

    if not input_dir.exists():
        raise FileNotFoundError(f"Input directory not found: {input_dir}")

    common_rules = build_rules(config.get("commonRules", []))
    file_rules_raw = config.get("fileRules", {})
    requested_files = config.get("files") or [path.name for path in sorted(input_dir.glob("*.xlsx"))]

    LOGGER.info("Job ID: %s", config.get("jobId", "unknown"))
    LOGGER.info("Files: %s", ", ".join(requested_files))

    for file_name in requested_files:
        source_path = input_dir / file_name
        if not source_path.exists():
            LOGGER.error("%s: input file not found", file_name)
            logs.append({"file": file_name, "stage": "input", "status": "failed", "message": "input file not found"})
            failures += 1
            continue

        try:
            LOGGER.info("%s: edit started", file_name)
            combined_rules = common_rules + build_rules(file_rules_raw.get(file_name, []))
            edited_path = edited_dir / file_name
            if combined_rules:
                apply_rules(source_path, edited_path, combined_rules)
            else:
                edited_path.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(source_path, edited_path)
            LOGGER.info("%s: edit completed", file_name)

            LOGGER.info("%s: PDF conversion started", file_name)
            pdf_path = convert_to_pdf(edited_path, pdf_dir)
            LOGGER.info("%s: PDF conversion completed: %s", file_name, pdf_path.name)
            logs.append({"file": file_name, "stage": "convert", "status": "success", "message": pdf_path.name})
        except Exception as exc:  # noqa: BLE001 - keep processing other files and write a clear job report.
            LOGGER.exception("%s: failed", file_name)
            logs.append({"file": file_name, "stage": "process", "status": "failed", "message": str(exc)})
            failures += 1

    report_path = output_dir / "conversion-report.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps({"failures": failures, "items": logs}, ensure_ascii=False, indent=2), encoding="utf-8")

    create_zip(pdf_dir, zip_path)
    LOGGER.info("ZIP created: %s", zip_path)
    return 1 if failures else 0


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    args = parse_args()
    try:
        return process_job(args.job_dir, args.output_dir, args.zip_path)
    except Exception:
        LOGGER.exception("Job failed before file processing completed")
        return 1


if __name__ == "__main__":
    sys.exit(main())
